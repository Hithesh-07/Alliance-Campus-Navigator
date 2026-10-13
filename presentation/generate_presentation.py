import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Presentation Setup ---
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# --- Theme Color Palette ---
COLOR_NAVY = RGBColor(15, 23, 42)       # Slate 900
COLOR_DARK_BLUE = RGBColor(30, 41, 59)   # Slate 800
COLOR_PRIMARY_TEAL = RGBColor(14, 116, 144) # Cyan 700
COLOR_ACCENT_TEAL = RGBColor(13, 148, 136) # Teal 600
COLOR_GREEN = RGBColor(22, 163, 74)     # Emerald / Green 600
COLOR_LIGHT_BG = RGBColor(248, 250, 252) # Slate 50
COLOR_CARD_BG = RGBColor(255, 255, 255)
COLOR_CARD_BORDER = RGBColor(226, 232, 240)
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)
COLOR_WHITE = RGBColor(255, 255, 255)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

SCREENSHOTS_DIR = os.path.join(os.path.dirname(__file__), "screenshots")

def add_header(slide, title_text, category_tag="ALLIANCE CAMPUS NAVIGATOR"):
    """Adds a standard modern slide header with tag and divider line."""
    # Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_tag.upper()
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_PRIMARY_TEAL

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY

    # Subtle divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_CARD_BORDER
    line.line.color.rgb = COLOR_CARD_BORDER

def add_footer(slide, slide_num):
    """Adds a clean academic footer with university, project info, and slide number."""
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(10.0), Inches(0.3))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Alliance Campus Navigator  |  B.Tech CSE Micro Project  |  Alliance University, Bengaluru"
    p.font.name = FONT_BODY
    p.font.size = Pt(9)
    p.font.color.rgb = COLOR_TEXT_MUTED

    num_box = slide.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.0), Inches(0.3))
    tf_num = num_box.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.text = f"{slide_num} / 12"
    p_num.font.name = FONT_BODY
    p_num.font.size = Pt(9)
    p_num.font.bold = True
    p_num.font.color.rgb = COLOR_PRIMARY_TEAL

def create_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    """Draws a clean rounded rectangle card container."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
# Dark Navy background
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = COLOR_NAVY
bg1.line.fill.background()

# Top accent bar
top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
top_bar.fill.solid()
top_bar.fill.fore_color.rgb = COLOR_ACCENT_TEAL
top_bar.line.fill.background()

# Title text frame (Left Side)
title_card = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(7.0), Inches(4.5))
tf1 = title_card.text_frame
tf1.word_wrap = True

p_badge = tf1.paragraphs[0]
p_badge.text = "ALLIANCE UNIVERSITY  •  DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING"
p_badge.font.name = FONT_HEADING
p_badge.font.size = Pt(11)
p_badge.font.bold = True
p_badge.font.color.rgb = COLOR_ACCENT_TEAL

p_main = tf1.add_paragraph()
p_main.space_before = Pt(14)
p_main.text = "Alliance Campus Navigator"
p_main.font.name = FONT_HEADING
p_main.font.size = Pt(36)
p_main.font.bold = True
p_main.font.color.rgb = COLOR_WHITE

p_sub = tf1.add_paragraph()
p_sub.space_before = Pt(8)
p_sub.text = "An Interactive Java-Based Campus Facility Finder"
p_sub.font.name = FONT_BODY
p_sub.font.size = Pt(18)
p_sub.font.color.rgb = RGBColor(203, 213, 225)

p_div = tf1.add_paragraph()
p_div.space_before = Pt(24)
p_div.text = "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━"
p_div.font.size = Pt(10)
p_div.font.color.rgb = COLOR_DARK_BLUE

p_meta = tf1.add_paragraph()
p_meta.space_before = Pt(14)
p_meta.text = "Student: B. Hithesh Reddy\nCourse: B.Tech in Computer Science and Engineering\nProject Type: Java Micro Project\nAcademic Year: 2026–2027"
p_meta.font.name = FONT_BODY
p_meta.font.size = Pt(13)
p_meta.font.color.rgb = RGBColor(241, 245, 249)

# Right Side: Application Screenshot Card
map_img_path = os.path.join(SCREENSHOTS_DIR, "2_campus_map_overview.png")
if os.path.exists(map_img_path):
    img_card = create_card(slide1, Inches(7.8), Inches(1.5), Inches(4.7), Inches(4.8), COLOR_DARK_BLUE, COLOR_PRIMARY_TEAL)
    slide1.shapes.add_picture(map_img_path, Inches(8.0), Inches(1.8), Inches(4.3), Inches(2.7))
    
    caption_box = slide1.shapes.add_textbox(Inches(8.0), Inches(4.7), Inches(4.3), Inches(1.3))
    tf_cap = caption_box.text_frame
    tf_cap.word_wrap = True
    p_cap = tf_cap.paragraphs[0]
    p_cap.text = "Live Desktop Application Screenshot\n• Java Swing & Java2D Vector Map Engine\n• Drag-and-Drop Facility Repositioning\n• Offline Persistent CSV Storage"
    p_cap.font.name = FONT_BODY
    p_cap.font.size = Pt(11)
    p_cap.font.color.rgb = RGBColor(226, 232, 240)

# ==============================================================================
# SLIDE 2: INTRODUCTION
# ==============================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide2, "Introduction — Project Background & Motivation", "OVERVIEW")
add_footer(slide2, 2)

# Card 1: Campus Reality
c1 = create_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
tb1 = slide2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.6))
tf = tb1.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🏛 The University Campus Environment"
p.font.name = FONT_HEADING
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

bullets1 = [
    "Alliance University encompasses a sprawling, multi-acre central campus in Anekal, Bengaluru.",
    "The campus houses numerous diverse departments: Academic schools, administrative blocks, student hostels, sports grounds, and food courts.",
    "Navigating between departments poses a challenge for freshmen, visiting parents, guest lecturers, and recruitment partners.",
    "Traditional physical signboards or static image maps fail to provide searchability, categorized filtering, or interactive facility details."
]
for b in bullets1:
    p = tf.add_paragraph()
    p.space_before = Pt(12)
    p.text = "• " + b
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MAIN

# Card 2: Campus Navigator Solution
c2 = create_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
tb2 = slide2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.6))
tf = tb2.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "💡 The Campus Navigator Solution"
p.font.name = FONT_HEADING
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY_TEAL

bullets2 = [
    "A standalone, offline-first Java desktop application built tailored specifically for Alliance University.",
    "Dynamic Java2D Vector Schematic: Illustrates core campus geography including roads, green spaces, and key building footprints.",
    "Structured Facility Database: Organizes 20 core university facilities into 6 functional categories.",
    "Unified User Experience: Offers instant multi-criteria search, one-click building centering, and real-time layout management.",
    "Zero External Dependencies: Runs purely on standard Java SE runtime with lightweight CSV file persistence."
]
for b in bullets2:
    p = tf.add_paragraph()
    p.space_before = Pt(12)
    p.text = "✔ " + b
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 3: PROBLEM STATEMENT & MOTIVATION
# ==============================================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide3, "Problem Statement & Project Motivation", "PROBLEM FORMULATION")
add_footer(slide3, 3)

# Problem Callout Box (Top)
prob_box = create_card(slide3, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.4), COLOR_DARK_BLUE, COLOR_PRIMARY_TEAL)
tb_prob = slide3.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.1), Inches(1.2))
tf = tb_prob.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "OFFICIAL PROBLEM STATEMENT"
p.font.name = FONT_HEADING
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_TEAL

p_body = tf.add_paragraph()
p_body.space_before = Pt(4)
p_body.text = '"To develop a user-friendly desktop application that presents campus facilities through an interactive schematic map and allows users to search, identify, and manage location information."'
p_body.font.name = FONT_HEADING
p_body.font.size = Pt(15)
p_body.font.bold = True
p_body.font.color.rgb = COLOR_WHITE

# 3 Column Cards below
cards_data = [
    ("⚠ Campus Navigation Challenges", [
        "Sprawling multi-block layout creates orientation difficulties.",
        "New students and visitors struggle to find lecture halls and admin offices.",
        "Static PDF/image campus maps lack keyword search and interaction."
    ], COLOR_CARD_BG),
    ("🎯 Target User Personas", [
        "First-Year Students: Quickly locate class blocks, labs, and dining.",
        "Visitors & Parents: Locate Administrative Block and seminar halls.",
        "Campus Admins: Add, edit, or reposition facilities dynamically."
    ], COLOR_CARD_BG),
    ("✨ Core Motivations", [
        "Eliminate dependency on proprietary, paid mapping APIs (Google Maps API).",
        "Deliver 100% offline reliability without internet access.",
        "Build a robust, extensible Java OOP showcase for B.Tech evaluation."
    ], COLOR_CARD_BG)
]

left_pos = 0.8
for title, items, bg in cards_data:
    create_card(slide3, Inches(left_pos), Inches(3.1), Inches(3.7), Inches(3.6), bg, COLOR_CARD_BORDER)
    tb = slide3.shapes.add_textbox(Inches(left_pos + 0.2), Inches(3.3), Inches(3.3), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_TEAL
    
    for item in items:
        p = tf.add_paragraph()
        p.space_before = Pt(10)
        p.text = "• " + item
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MAIN
    left_pos += 4.0

# ==============================================================================
# SLIDE 4: OBJECTIVES & SCOPE
# ==============================================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide4, "Project Objectives & Scope Boundaries", "PROJECT SCOPE")
add_footer(slide4, 4)

# Objectives (Left Column)
create_card(slide4, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.1))
tb_obj = slide4.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.4), Inches(4.7))
tf = tb_obj.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🎯 Core Project Objectives"
p.font.name = FONT_HEADING
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

objs = [
    ("Desktop GUI Engineering", "Build a high-performance desktop application using modern Java SE (Swing & AWT)."),
    ("Dynamic Vector Map", "Draw an illustrative schematic map using Java2D vector primitives—roads, grounds, and buildings."),
    ("Category Taxonomy", "Organize 20 campus landmarks into 6 color-coded operational categories."),
    ("Multi-Criteria Search", "Enable instant search across facility names, categories, and keyword descriptions."),
    ("Interactive Dragging", "Allow dragging buildings on the canvas with coordinate persistence."),
    ("File Persistence", "Manage records using UTF-8 RFC 4180 CSV files without database setup.")
]
for title, desc in objs:
    p = tf.add_paragraph()
    p.space_before = Pt(8)
    p.text = f"✔ {title}: "
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_TEAL
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = COLOR_TEXT_MAIN

# Scope & Limitations (Right Column)
create_card(slide4, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.1))
tb_scp = slide4.shapes.add_textbox(Inches(7.4), Inches(1.8), Inches(4.8), Inches(4.7))
tf = tb_scp.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "⚖ Scope & Boundary Definitions"
p.font.name = FONT_HEADING
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_MAIN

scopes = [
    ("Schematic Diagram", "The map represents visual spatial relationships rather than sub-meter satellite survey accuracy."),
    ("No GPS/Turn-by-Turn", "Intentionally designed for orientation; does not require GPS hardware, mobile sensors, or turn-by-turn routing."),
    ("Self-Contained Data", "Accuracy and positions are derived from the local CSV dataset (locations.csv & map_layout.csv)."),
    ("Local Persistence", "Stores layout and location changes locally on the host machine without cloud dependencies.")
]
for title, desc in scopes:
    p = tf.add_paragraph()
    p.space_before = Pt(12)
    p.text = f"• {title}:\n  "
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = COLOR_TEXT_MUTED

# ==============================================================================
# SLIDE 5: TECHNOLOGY STACK
# ==============================================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide5, "Technology Stack & Architectural Tooling", "TECHNOLOGY SELECTION")
add_footer(slide5, 5)

tech_cards = [
    ("☕ Java 21+ Platform", "Core Programming Language", [
        "Provides modern object-oriented language features and high performance.",
        "Utilizes standard Java SE runtime with zero external JAR library dependencies.",
        "Ensures cross-platform desktop execution on Windows, macOS, and Linux."
    ]),
    ("🖥 Java Swing & AWT", "Desktop Graphical User Interface", [
        "Manages application shell, sidebar, card layouts, dialogs, and tables.",
        "Guarantees thread-safe UI updates on Swing's Event Dispatch Thread (EDT).",
        "Uses custom component painting for modern rounded cards, buttons, and badges."
    ]),
    ("🎨 Java2D Graphics API", "Custom Vector Map Engine", [
        "Dynamically renders roads, pedestrian paths, green spaces, and buildings.",
        "Applies anti-aliasing rendering hints for clean visual presentation.",
        "Executes world-to-screen AffineTransform transformations for smooth zoom & pan."
    ]),
    ("📄 RFC 4180 CSV Storage", "Persistence Layer", [
        "Implements robust UTF-8 CSV parsing and writing via CsvUtils.",
        "Safely escapes embedded commas, quotes, and multi-line descriptions.",
        "Eliminates SQL database setup, ODBC drivers, and administrative server overhead."
    ]),
    ("💻 VS Code & PowerShell", "Development & Build Environment", [
        "Organized project structure with clean separation of src, data, out, and test.",
        "Direct compilation via javac -encoding UTF-8 -d out.",
        "One-click launch script run.ps1 for seamless evaluation and live demonstration."
    ]),
    ("🐙 Git & GitHub", "Version Control & Collaboration", [
        "Full version control tracking code evolution and file revisions.",
        "Public GitHub repository: Hithesh-07/Alliance-Campus-Navigator.",
        "Transparent commit history and automated unit/integration test coverage."
    ])
]

# 2 rows x 3 columns
positions = [
    (0.8, 1.6), (4.8, 1.6), (8.8, 1.6),
    (0.8, 4.3), (4.8, 4.3), (8.8, 4.3)
]

for idx, (title, sub, points) in enumerate(tech_cards):
    lx, ty = positions[idx]
    create_card(slide5, Inches(lx), Inches(ty), Inches(3.7), Inches(2.5))
    tb = slide5.shapes.add_textbox(Inches(lx + 0.15), Inches(ty + 0.15), Inches(3.4), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_TEAL

    p_sub = tf.add_paragraph()
    p_sub.text = sub
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    
    for pt in points:
        p = tf.add_paragraph()
        p.space_before = Pt(4)
        p.text = "• " + pt
        p.font.name = FONT_BODY
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 6: SYSTEM ARCHITECTURE
# ==============================================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide6, "System Architecture & End-to-End Data Flow", "ARCHITECTURE")
add_footer(slide6, 6)

# Architecture Flow Diagram using Cards
# Layer 1: Presentation Layer
create_card(slide6, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.1), COLOR_DARK_BLUE, COLOR_PRIMARY_TEAL)
tb_l1 = slide6.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(0.9))
tf = tb_l1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "PRESENTATION LAYER (ui package)"
p.font.name = FONT_HEADING
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_TEAL
p_desc = tf.add_paragraph()
p_desc.text = "CampusNavFrame (Shell & Nav)  ◄►  DashboardPanel  |  CampusMapPanel (Java2D Canvas)  |  LocationSearchPanel  |  LocationManagementPanel"
p_desc.font.name = FONT_BODY
p_desc.font.size = Pt(12)
p_desc.font.bold = True
p_desc.font.color.rgb = COLOR_WHITE

# Arrow down
arr1 = slide6.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.4), Inches(2.8), Inches(0.5), Inches(0.3))
arr1.fill.solid()
arr1.fill.fore_color.rgb = COLOR_ACCENT_TEAL
arr1.line.fill.background()

# Layer 2: Business Logic & Access Layer
create_card(slide6, Inches(0.8), Inches(3.2), Inches(7.5), Inches(1.8))
tb_l2 = slide6.shapes.add_textbox(Inches(1.0), Inches(3.3), Inches(7.1), Inches(1.6))
tf = tb_l2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "APPLICATION & BUSINESS LOGIC LAYER (data package)"
p.font.name = FONT_HEADING
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY_TEAL

bullets_l2 = [
    "CampusDataManager (Singleton): Central state controller & observable data source.",
    "DataChangeListener: Reactive observer notifying UI views of additions, updates, or deletions.",
    "DefaultCampusData: Hardcoded baseline fallback for instant initialization and Reset Layout.",
    "CsvUtils: RFC 4180 parsing engine with quote escaping and stream buffering."
]
for b in bullets_l2:
    p = tf.add_paragraph()
    p.space_before = Pt(3)
    p.text = "• " + b
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN

# Model Box (Right of Layer 2)
create_card(slide6, Inches(8.6), Inches(3.2), Inches(3.9), Inches(1.8))
tb_model = slide6.shapes.add_textbox(Inches(8.8), Inches(3.3), Inches(3.5), Inches(1.6))
tf = tb_model.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "DATA DOMAIN MODELS (model package)"
p.font.name = FONT_HEADING
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_GREEN

models = [
    "CampusLocation: id, name, category, description",
    "MapItem: id, x, y, width, height, shapeType",
    "Category Enum: ACADEMIC, ADMINISTRATIVE, FOOD, SPORTS, RESIDENTIAL, OTHER"
]
for m in models:
    p = tf.add_paragraph()
    p.space_before = Pt(4)
    p.text = "• " + m
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN

# Arrow down to Persistence
arr2 = slide6.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(4.3), Inches(5.1), Inches(0.5), Inches(0.3))
arr2.fill.solid()
arr2.fill.fore_color.rgb = COLOR_ACCENT_TEAL
arr2.line.fill.background()

# Layer 3: Persistence Layer
create_card(slide6, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.3), COLOR_LIGHT_BG, COLOR_PRIMARY_TEAL)
tb_l3 = slide6.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.3), Inches(1.1))
tf = tb_l3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "PERSISTENCE STORAGE LAYER (data/ directory)"
p.font.name = FONT_HEADING
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY_TEAL

p_csv = tf.add_paragraph()
p_csv.text = "• data/locations.csv ── Stores metadata (id, name, category, description)\n• data/map_layout.csv ── Stores 2D spatial footprints (id, x, y, width, height, shapeType, customColorHex)"
p_csv.font.name = FONT_BODY
p_csv.font.size = Pt(12)
p_csv.font.bold = True
p_csv.font.color.rgb = COLOR_NAVY

# ==============================================================================
# SLIDE 7: CAMPUS MAP & FACILITY CATEGORIES
# ==============================================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide7, "Campus Schematic Map & Facility Taxonomies", "CAMPUS GEOGRAPHY")
add_footer(slide7, 7)

# Left: Real Map Screenshot
map_screen = os.path.join(SCREENSHOTS_DIR, "3_campus_map_selected.png")
if os.path.exists(map_screen):
    create_card(slide7, Inches(0.8), Inches(1.6), Inches(6.6), Inches(5.1), COLOR_CARD_BG, COLOR_PRIMARY_TEAL)
    slide7.shapes.add_picture(map_screen, Inches(0.9), Inches(1.7), Inches(6.4), Inches(4.1))
    
    cap_box = slide7.shapes.add_textbox(Inches(0.9), Inches(5.9), Inches(6.4), Inches(0.7))
    tf = cap_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Actual Java2D vector rendering showing selected Alliance School of Business (Rotunda), Library Dome, Football Ground, and Cricket Oval."
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MUTED

# Right: Categories & Featured Locations
create_card(slide7, Inches(7.7), Inches(1.6), Inches(4.8), Inches(5.1))
tb_cat = slide7.shapes.add_textbox(Inches(7.9), Inches(1.7), Inches(4.4), Inches(4.8))
tf = tb_cat.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🏛 Campus Facility Categorization"
p.font.name = FONT_HEADING
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

cat_data = [
    ("Academic (Blue)", "Alliance School of Business, Library, Eng. Block, Law, Liberal Arts, LC-1"),
    ("Administrative (Purple)", "Administrative Block (Chancellor's office, Admissions, Registrar)"),
    ("Food & Dining (Orange)", "Alliance Food Court, Daana Paani Multicuisine Cafe"),
    ("Sports & Recreation (Green)", "Football Turf Ground, Cricket Ground, Volleyball, Basketball"),
    ("Residential (Amber)", "Hall of Residence 1 (HOR 1), HOR 2, HOR 5, Staff Quarters"),
    ("Other Services (Slate)", "Water Treatment & Recycling Plant, Western Gateway")
]
for cat, locs in cat_data:
    p = tf.add_paragraph()
    p.space_before = Pt(8)
    p.text = f"• {cat}:\n  "
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_TEAL
    run = p.add_run()
    run.text = locs
    run.font.bold = False
    run.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 8: USER INTERFACE & CORE FEATURES
# ==============================================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide8, "User Interface Showcase & Interactive Features", "USER EXPERIENCE")
add_footer(slide8, 8)

# 3 Screenshots side by side with feature labels
ui_items = [
    ("Dashboard Analytics View", "1_dashboard_view.png", [
        "Live summary metric counters (Academic, Food, Sports, Hostels).",
        "Interactive mini-map preview card with instant launch shortcut.",
        "Quick-access popular university hubs list."
    ]),
    ("Multi-Criteria Search Engine", "4_search_view.png", [
        "Real-time keystroke filtering across names, IDs, and descriptions.",
        "One-click category filter chips (All, Academic, Sports, etc.).",
        "'Show on Map' button automatically navigates & highlights."
    ]),
    ("Location Management (CRUD)", "5_management_view.png", [
        "Searchable table showing coordinates (X, Y) and dimensions.",
        "Modal dialogs for Add / Edit with validation checks.",
        "Instant cross-view sync and persistent CSV updates."
    ])
]

left_ui = 0.8
for title, img_name, highlights in ui_items:
    create_card(slide8, Inches(left_ui), Inches(1.6), Inches(3.7), Inches(5.1))
    
    # Title
    tb = slide8.shapes.add_textbox(Inches(left_ui + 0.15), Inches(1.7), Inches(3.4), Inches(0.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY
    
    # Image
    img_p = os.path.join(SCREENSHOTS_DIR, img_name)
    if os.path.exists(img_p):
        slide8.shapes.add_picture(img_p, Inches(left_ui + 0.15), Inches(2.15), Inches(3.4), Inches(2.1))
    
    # Bullets
    tb_b = slide8.shapes.add_textbox(Inches(left_ui + 0.15), Inches(4.35), Inches(3.4), Inches(2.2))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    for h in highlights:
        p = tf_b.add_paragraph()
        p.space_before = Pt(3)
        p.text = "• " + h
        p.font.name = FONT_BODY
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MAIN
        
    left_ui += 4.0

# ==============================================================================
# SLIDE 9: JAVA IMPLEMENTATION & OOP PRINCIPLES
# ==============================================================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide9, "Object-Oriented Programming (OOP) Implementation", "SOURCE CODE DESIGN")
add_footer(slide9, 9)

# Left Column: OOP Pillars in Project
create_card(slide9, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.1))
tb_oop = slide9.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.4), Inches(4.8))
tf = tb_oop.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🏛 OOP Principles Applied in Source Code"
p.font.name = FONT_HEADING
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

oops = [
    ("Inheritance", "CampusNavFrame extends JFrame; CampusMapPanel, DashboardPanel, LocationSearchPanel extend JPanel; LocationDialog extends JDialog."),
    ("Polymorphism", "@Override paintComponent(Graphics g) customized across map and buttons; interface implementation in NavigationController."),
    ("Encapsulation", "CampusLocation and MapItem enforce private fields with strict getters, setters, and boundary constraints (e.g. Math.max(30, width))."),
    ("Event-Driven Architecture", "Swing event listeners (MouseAdapter, KeyAdapter) handle real-time canvas dragging and live search filtering."),
    ("Observer Pattern", "DataChangeListener decoupled event bus triggers instant cross-panel refreshes when records change.")
]
for title, desc in oops:
    p = tf.add_paragraph()
    p.space_before = Pt(8)
    p.text = f"• {title}: "
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_TEAL
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = COLOR_TEXT_MAIN

# Right Column: Actual Source Code Snippet
create_card(slide9, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.1), COLOR_DARK_BLUE, COLOR_PRIMARY_TEAL)
tb_code = slide9.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.2), Inches(4.8))
tf = tb_code.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "💻 Actual Source Code Example: Map Drag Translation"
p.font.name = FONT_HEADING
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_TEAL

code_snippet = """// From ui/CampusMapPanel.java (World Coordinate Translation)
public Point2D toWorld(Point screenPt) {
    double wx = (screenPt.x - panX) / zoom;
    double wy = (screenPt.y - panY) / zoom;
    return new Point2D.Double(wx, wy);
}

@Override
public void mouseDragged(MouseEvent e) {
    if (isDraggingBuilding && draggingItem != null) {
        Point2D worldPt = toWorld(e.getPoint());
        double newX = worldPt.getX() + dragStartBuildingX;
        double newY = worldPt.getY() + dragStartBuildingY;
        
        draggingItem.setX(Math.round(newX));
        draggingItem.setY(Math.round(newY));
        hasUnsavedChanges = true;
        repaint();
    }
}"""

p_code = tf.add_paragraph()
p_code.space_before = Pt(8)
p_code.text = code_snippet
p_code.font.name = "Consolas"
p_code.font.size = Pt(9.5)
p_code.font.color.rgb = RGBColor(226, 232, 240)

p_expl = tf.add_paragraph()
p_expl.space_before = Pt(8)
p_expl.text = "Key Logic: Screen pixels are inverted through zoom/pan transforms, ensuring that dragging never corrupts saved logical coordinates."
p_expl.font.name = FONT_BODY
p_expl.font.size = Pt(10)
p_expl.font.color.rgb = COLOR_ACCENT_TEAL

# ==============================================================================
# SLIDE 10: FILE HANDLING & DATA PERSISTENCE
# ==============================================================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide10, "File Handling & Persistent CSV Architecture", "DATA PERSISTENCE")
add_footer(slide10, 10)

# Table showing CSV File Specifications
table_shape = slide10.shapes.add_table(3, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.2))
tbl = table_shape.table
tbl.columns[0].width = Inches(2.4)
tbl.columns[1].width = Inches(3.3)
tbl.columns[2].width = Inches(3.6)
tbl.columns[3].width = Inches(2.4)

headers = ["Persistent File", "CSV Column Schema", "Domain Purpose", "Read / Write Lifecycle"]
for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_NAVY
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

rows_data = [
    ("data/locations.csv", "id, name, category, description", "Stores facility identities, names, category groupings, and editable descriptions.", "Loaded on start-up; written on Add, Edit, or Delete."),
    ("data/map_layout.csv", "id, x, y, width, height, shapeType, customColorHex", "Stores 2D spatial positions, bounding dimensions, and architectural shape geometry.", "Loaded on start-up; written on Save Layout & reset on Reset Layout.")
]
for r_idx, r_vals in enumerate(rows_data):
    for c_idx, val in enumerate(r_vals):
        cell = tbl.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_BODY
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT_MAIN

# Lower Section: Key Persistence Mechanisms
create_card(slide10, Inches(0.8), Inches(4.1), Inches(11.7), Inches(2.6))
tb_mech = slide10.shapes.add_textbox(Inches(1.1), Inches(4.2), Inches(11.1), Inches(2.3))
tf = tb_mech.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "⚡ Key Implementation Highlights in CsvUtils & CampusDataManager"
p.font.name = FONT_HEADING
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY_TEAL

mechs = [
    "RFC 4180 Compliance: Handles quotes, commas inside text, and escaped quotes (\"\") through custom state-machine parsing.",
    "UTF-8 Encoding: Guarantees safe encoding across all operating systems without character corruption.",
    "Automatic Initial Seeding: If CSV files are missing on first run, DefaultCampusData automatically populates all 20 records.",
    "Data Integrity: Synchronizes memory models with disk; deleting a location removes both its record and its spatial map footprint."
]
for m in mechs:
    p = tf.add_paragraph()
    p.space_before = Pt(6)
    p.text = "✔ " + m
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 11: TESTING, RESULTS & LIMITATIONS
# ==============================================================================
slide11 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide11, "Verification Testing, Results & Known Limitations", "QUALITY ASSURANCE")
add_footer(slide11, 11)

# Test Results Table
table_shape2 = slide11.shapes.add_table(6, 4, Inches(0.8), Inches(1.6), Inches(7.2), Inches(5.1))
tbl2 = table_shape2.table
tbl2.columns[0].width = Inches(1.8)
tbl2.columns[1].width = Inches(2.3)
tbl2.columns[2].width = Inches(2.1)
tbl2.columns[3].width = Inches(1.0)

t_headers = ["Test Suite", "Test Objective", "Verification Method", "Status"]
for i, h in enumerate(t_headers):
    cell = tbl2.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_NAVY
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

tests_exec = [
    ("CSV Escaping", "RFC 4180 quotes & commas", "testCsvEscaping()", "PASSED"),
    ("Default Seeding", "20 default campus records", "testDefaultDataIntegrity()", "PASSED"),
    ("CRUD Lifecycle", "Add, edit, delete & sync", "testDataManagerPersistence()", "PASSED"),
    ("Search & Filter", "Case-insensitive queries", "testSearchFiltering()", "PASSED"),
    ("Map Hit Bounds", "World-space coordinate hit", "testMapGeometryAndBounds()", "PASSED")
]
for r_idx, r_vals in enumerate(tests_exec):
    for c_idx, val in enumerate(r_vals):
        cell = tbl2.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_BODY
        p.font.size = Pt(9.5)
        if c_idx == 3:
            p.font.bold = True
            p.font.color.rgb = COLOR_GREEN
        else:
            p.font.color.rgb = COLOR_TEXT_MAIN

# Right: GUI Smoke Test & Limitations
create_card(slide11, Inches(8.3), Inches(1.6), Inches(4.2), Inches(5.1))
tb_lim = slide11.shapes.add_textbox(Inches(8.5), Inches(1.7), Inches(3.8), Inches(4.8))
tf = tb_lim.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🖥 GUI Smoke Test Verification"
p.font.name = FONT_HEADING
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY_TEAL

p_smoke = tf.add_paragraph()
p_smoke.space_before = Pt(4)
p_smoke.text = "CampusNavGuiSmokeTest verified frame creation on Swing EDT, view card switching, and layout resizing without headless or rendering exceptions."
p_smoke.font.name = FONT_BODY
p_smoke.font.size = Pt(10)
p_smoke.font.color.rgb = COLOR_TEXT_MAIN

p_lim_t = tf.add_paragraph()
p_lim_t.space_before = Pt(14)
p_lim_t.text = "⚖ Documented Limitations"
p_lim_t.font.name = FONT_HEADING
p_lim_t.font.size = Pt(13)
p_lim_t.font.bold = True
p_lim_t.font.color.rgb = RGBColor(225, 29, 72)

lims = [
    "Schematic map is an illustrative diagram, not a GPS GIS projection.",
    "No turn-by-turn routing (intentionally out of scope).",
    "Single-user desktop architecture without multi-user concurrent locking."
]
for l in lims:
    p = tf.add_paragraph()
    p.space_before = Pt(4)
    p.text = "• " + l
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MUTED

# ==============================================================================
# SLIDE 12: CONCLUSION & FUTURE SCOPE
# ==============================================================================
slide12 = prs.slides.add_slide(prs.slide_layouts[6])
# Dark Navy background
bg12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg12.fill.solid()
bg12.fill.fore_color.rgb = COLOR_NAVY
bg12.line.fill.background()

# Top accent bar
top_bar12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
top_bar12.fill.solid()
top_bar12.fill.fore_color.rgb = COLOR_ACCENT_TEAL
top_bar12.line.fill.background()

# Header on dark
h_box = slide12.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
tf = h_box.text_frame
p = tf.paragraphs[0]
p.text = "Conclusion, Future Scope & Viva Discussion"
p.font.name = FONT_HEADING
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE

# 2 Cards on Dark
# Card 1: Project Achievements
c_ach = create_card(slide12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.3), COLOR_DARK_BLUE, COLOR_PRIMARY_TEAL)
tb_ach = slide12.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(3.8))
tf = tb_ach.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🏆 Project Summary & Achievements"
p.font.name = FONT_HEADING
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_TEAL

achs = [
    "Successfully developed an offline-first desktop navigation system for Alliance University.",
    "Integrated Java2D custom drawing with real-time drag-and-drop map layout modification.",
    "Implemented full CRUD lifecycle with synchronized CSV storage.",
    "Demonstrated sound OOP architecture, clean separation of concerns, and thread-safe Swing execution.",
    "100% test pass rate across backend persistence and GUI smoke suites."
]
for a in achs:
    p = tf.add_paragraph()
    p.space_before = Pt(8)
    p.text = "✔ " + a
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_WHITE

# Card 2: Future Scope
c_fut = create_card(slide12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.3), COLOR_DARK_BLUE, COLOR_PRIMARY_TEAL)
tb_fut = slide12.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(3.8))
tf = tb_fut.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🚀 Future Enhancement Scope"
p.font.name = FONT_HEADING
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_TEAL

futs = [
    "Shortest Path Routing: Implement Dijkstra or A* pathfinding algorithm for pedestrian pathways.",
    "Database Integration: Migrate from CSV to SQLite/PostgreSQL for enterprise scalability.",
    "Mobile & Web Companion: Develop responsive companion apps using modern web technologies.",
    "Multi-Floor Indoor Navigation: Support level-by-level indoor layout diagrams for complex academic blocks."
]
for f in futs:
    p = tf.add_paragraph()
    p.space_before = Pt(8)
    p.text = "• " + f
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_WHITE

# Bottom Banner on Slide 12
bot_card = create_card(slide12, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9), COLOR_NAVY, COLOR_ACCENT_TEAL)
tb_bot = slide12.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.8))
tf = tb_bot.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🔗 GitHub Repository: https://github.com/Hithesh-07/Alliance-Campus-Navigator"
p.font.name = FONT_HEADING
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_TEAL

p_ty = tf.add_paragraph()
p_ty.space_before = Pt(2)
p_ty.text = "Thank You! Questions & Viva Discussion Welcome."
p_ty.font.name = FONT_HEADING
p_ty.font.size = Pt(13)
p_ty.font.bold = True
p_ty.font.color.rgb = COLOR_WHITE

# --- Save Presentation ---
output_pptx = os.path.join(os.path.dirname(__file__), "Alliance_Campus_Navigator_Presentation.pptx")
prs.save(output_pptx)
print(f"Presentation successfully generated at: {output_pptx}")
